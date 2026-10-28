"""
CLI Script: Generate Benchmark Metrics and Evaluation Report.

Usage:
  python scripts/run_evaluation.py --input "output/ai_tagged_zenvia_curated_men_tshirts.xlsx"
"""

import sys
import io
import os
import argparse
import json
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from zenvia_ai.config import Config, setup_logger
from zenvia_ai.models import ProductItem, ClassificationResult
from zenvia_ai.evaluator import ModelEvaluator

logger = setup_logger("run_evaluation")


def evaluate_existing_workbook(excel_path: str) -> None:
    if not os.path.exists(excel_path):
        logger.error(f"Spreadsheet not found: {excel_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    header_map = {name: idx for idx, name in enumerate(headers)}

    pairs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        def val(col_name: str) -> str:
            idx = header_map.get(col_name)
            if idx is not None and idx < len(row) and row[idx] is not None:
                return str(row[idx]).strip()
            return ""

        def num(col_name: str, default: float = 0.0) -> float:
            try:
                return float(val(col_name))
            except ValueError:
                return default

        name = val("name")
        ai_vibes = val("ai_vibeTags")
        if not name or not ai_vibes:
            continue

        item = ProductItem(
            name=name,
            category=val("category") or "men",
            manual_vibes=val("vibeTags"),
            manual_subcategory=val("subcategory"),
            main_image=val("mainImage"),
        )

        res = ClassificationResult(
            vibe_tags=[v.strip() for v in val("ai_vibeTags").split(",") if v.strip()],
            subcategory=val("ai_subcategory"),
            occasion=val("ai_occasion") or "casual",
            confidence=num("ai_confidence", 0.90),
            visual_note=val("ai_visualNote"),
            latency_seconds=num("ai_latency_sec", 1.2),
            status=val("ai_status") or "ok",
        )
        pairs.append((item, res))

    wb.close()

    metrics = ModelEvaluator.evaluate(pairs)
    markdown_report = ModelEvaluator.generate_markdown_report(metrics)

    report_path = os.path.join(Config.OUTPUT_DIR, "evaluation_report.md")
    metrics_path = os.path.join(Config.OUTPUT_DIR, "evaluation_metrics.json")

    with open(report_path, "w", encoding="utf-8") as f:
        f.write(markdown_report)

    with open(metrics_path, "w", encoding="utf-8") as f:
        json.dump(metrics.model_dump(), f, indent=2)

    print("\n" + markdown_report)
    print("\n" + "=" * 65)
    print(f"📄 Markdown Report Saved: {report_path}")
    print(f"📊 Metrics JSON Saved:    {metrics_path}")
    print("=" * 65)


def main():
    parser = argparse.ArgumentParser(description="Zenvia AI Evaluation Benchmark CLI")
    parser.add_argument("--input", "-i", default=os.path.join(Config.OUTPUT_DIR, "ai_tagged_zenvia_curated_men_tshirts.xlsx"), help="Path to tagged Excel spreadsheet")
    args = parser.parse_args()
    evaluate_existing_workbook(args.input)


if __name__ == "__main__":
    main()
