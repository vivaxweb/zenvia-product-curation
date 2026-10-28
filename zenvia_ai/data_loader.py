"""
Catalog Ingestion and Serialization Utilities for Excel and Spreadsheets.
"""

import os
from typing import List, Generator, Tuple, Optional
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook import Workbook

from .models import ProductItem, ClassificationResult
from .config import setup_logger

logger = setup_logger("zenvia_ai.data_loader")


class CatalogLoader:
    """Manages reading raw product sheets and writing enriched outputs."""

    ENRICHED_COLUMNS: List[str] = [
        "ai_vibeTags",
        "ai_subcategory",
        "ai_occasion",
        "ai_confidence",
        "ai_visualNote",
        "ai_status",
        "ai_latency_sec",
    ]

    @staticmethod
    def load_products_from_excel(file_path: str, limit: Optional[int] = None) -> List[ProductItem]:
        """Load and parse products from an Excel workbook into ProductItem models."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Catalog file not found: {file_path}")

        wb: Workbook = openpyxl.load_workbook(file_path, data_only=True)
        ws: Worksheet = wb.active

        headers = [str(cell.value or "").strip() for cell in ws[1]]
        header_map = {name: idx for idx, name in enumerate(headers)}

        products: List[ProductItem] = []
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        total_rows = len(rows) if limit is None else min(len(rows), limit)

        for row in rows[:total_rows]:
            def val(col_name: str, default: str = "") -> str:
                idx = header_map.get(col_name)
                if idx is not None and idx < len(row) and row[idx] is not None:
                    return str(row[idx]).strip()
                return default

            def num(col_name: str, default: float = 0.0) -> float:
                raw = val(col_name)
                try:
                    return float(raw)
                except ValueError:
                    return default

            name = val("name")
            if not name:
                continue

            item = ProductItem(
                name=name,
                price=num("price"),
                discounted_price=num("discountedPrice"),
                description=val("description"),
                category=val("category", "men"),
                main_image=val("mainImage"),
                manual_vibes=val("vibeTags"),
                manual_subcategory=val("subcategory"),
                source=val("source", "unknown"),
                redirect_link=val("redirectLink"),
                rating=num("ratings"),
            )
            products.append(item)

        wb.close()
        logger.info(f"Loaded {len(products)} products from {os.path.basename(file_path)}")
        return products

    @classmethod
    def save_enriched_excel(
        cls,
        source_excel: str,
        results: List[Tuple[ProductItem, ClassificationResult]],
        output_path: str,
    ) -> str:
        """
        Write enriched AI predictions back into the Excel spreadsheet preserving all existing data.
        """
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb: Workbook = openpyxl.load_workbook(source_excel)
        ws: Worksheet = wb.active

        headers = [cell.value for cell in ws[1]]
        
        # Ensure enriched columns are registered in header row
        for col_name in cls.ENRICHED_COLUMNS:
            if col_name not in headers:
                headers.append(col_name)
                ws.cell(row=1, column=len(headers), value=col_name)

        col_index_map = {name: idx + 1 for idx, name in enumerate(headers)}

        # Update matching rows
        for row_idx, (item, res) in enumerate(results, start=2):
            if row_idx > ws.max_row:
                break
            ws.cell(row=row_idx, column=col_index_map["ai_vibeTags"], value=res.vibe_string)
            ws.cell(row=row_idx, column=col_index_map["ai_subcategory"], value=res.subcategory)
            ws.cell(row=row_idx, column=col_index_map["ai_occasion"], value=res.occasion)
            ws.cell(row=row_idx, column=col_index_map["ai_confidence"], value=res.confidence)
            ws.cell(row=row_idx, column=col_index_map["ai_visualNote"], value=res.visual_note)
            ws.cell(row=row_idx, column=col_index_map["ai_status"], value=res.status)
            ws.cell(row=row_idx, column=col_index_map["ai_latency_sec"], value=res.latency_seconds)

        wb.save(output_path)
        wb.close()
        logger.info(f"Enriched workbook saved successfully: {output_path}")
        return output_path
